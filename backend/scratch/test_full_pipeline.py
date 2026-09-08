import sys
sys.path.insert(0, '.')
from app.core.database import SessionLocal
from app.models.dataset import Dataset, DatasetColumn
from app.services.ai_service import AIService
from app.services.decision_engine import FormulaDecisionEngine
from app.services.formula_generator import FormulaGenerator
from app.services.execution_engine import ExecutionEngine
from app.services.report_service import ReportService
import openpyxl

db = SessionLocal()
d = db.query(Dataset).filter_by(id='81189e58-e059-4aba-bde9-2ac5c10e74cd').first()
cols = db.query(DatasetColumn).filter_by(dataset_id=d.id).all()

query = "Total penjualan Toyota di Pekanbaru pada Februari 2024"

# 1. Parse query
parsed = AIService.parse_query(query, cols)
print("Intent:", parsed.intent, parsed.operation, parsed.target_field)
print("Filters:")
for f in parsed.filters:
    print(f"  {f.field} {f.operator} {f.value}")

# 2. Decision Engine
engine = FormulaDecisionEngine()
decision = engine.decide_formula(parsed, cols)
print("Formula:", decision.formula_id, "Confidence:", decision.ml_confidence)
print("Explanation:", decision.reason)

# 3. Formula Generator
formula_str = FormulaGenerator.generate_excel_formula(decision, parsed, cols)
print("Generated Formula:", formula_str)

import pandas as pd
df = pd.read_excel(d.file_path, sheet_name="Data_Penjualan")
# 4. Execution Engine
summary, headers, rows, chart = ExecutionEngine.execute(df=df, intent=parsed, decision=decision)
print(f"Summary: {summary}")
print(f"Headers count: {len(headers)}")
print(f"Rows count: {len(rows)}")

# 5. Export Report
analysis_data = {
    "user_query": query,
    "formula_id": decision.formula_id,
    "formula_name": decision.formula_name,
    "generated_excel_formula": formula_str,
    "formula_explanation": decision.reason,
    "table_headers": headers,
    "table_rows": rows,
    "calculation_summary": summary,
}
export_path = ReportService.export_analysis_to_excel(analysis_data, "Test_Toyota_Pekanbaru")
print(f"Exported to: {export_path}")

# Check workbook
wb = openpyxl.load_workbook(export_path, data_only=False)
ws = wb.active
print("Title cell A1:", ws["A1"].value)
print("Formula label A5:", ws["A5"].value)
print("Formula value B5:", repr(ws["B5"].value), "data_type:", ws["B5"].data_type)
print("Table Header row 9:", [ws.cell(9, c).value for c in range(1, 8)])
print("Data row 10:", [ws.cell(10, c).value for c in range(1, 8)])
print("Last row:", ws.max_row, [ws.cell(ws.max_row, c).value for c in [1, 16]])
