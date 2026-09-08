import requests
import openpyxl

BASE_URL = "http://127.0.0.1:8000/api/v1"

# 1. Analyze query
payload = {
    "dataset_id": "81189e58-e059-4aba-bde9-2ac5c10e74cd",
    "query": "Total penjualan Toyota di Pekanbaru pada Februari 2024"
}
resp = requests.post(f"{BASE_URL}/query/analyze", json=payload)
data = resp.json()

print("Status Code:", resp.status_code)
analysis_id = data.get("analysis_id")
print("Analysis ID:", analysis_id)
print("Formula:", data.get("decision", {}).get("formula_id"))
print("Formula Excel:", data.get("decision", {}).get("generated_excel_formula"))
print("Table Rows Count:", len(data.get("table_rows", [])))
print("Summary:", data.get("calculation_summary"))

# 2. Export Excel
export_resp = requests.post(f"{BASE_URL}/reports/export/{analysis_id}")
export_data = export_resp.json()
print("Export Response:", export_data)
download_url = export_data.get("download_url")

# 3. Download & inspect Excel
dl_resp = requests.get(f"http://127.0.0.1:8000{download_url}")
with open("scratch/test_downloaded.xlsx", "wb") as f:
    f.write(dl_resp.content)

wb = openpyxl.load_workbook("scratch/test_downloaded.xlsx", data_only=False)
ws = wb.active
print("\n--- EXCEL INSPECTION ---")
print("Cell B3 (Query):", ws["B3"].value)
print("Cell B4 (Formula):", ws["B4"].value)
print("Cell B5 (Excel Formula):", ws["B5"].value, "type:", ws["B5"].data_type)
print("Cell B6 (Explanation):", ws["B6"].value)
print("Row 9 (Headers, first 7):", [ws.cell(9, c).value for c in range(1, 8)])
print("Row 10 (First data row):", [ws.cell(10, c).value for c in range(1, 8)])
print("Max Row:", ws.max_row)
print(f"Row {ws.max_row} (Total Row):", ws.cell(ws.max_row, 1).value, "Sum cell P:", ws.cell(ws.max_row, 16).value)
