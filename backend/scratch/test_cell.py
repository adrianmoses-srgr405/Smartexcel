import openpyxl
from pathlib import Path

wb = openpyxl.Workbook()
ws = wb.active
ws['A1'].value = "=SUMIFS(P:P, B:B)"
ws['A1'].data_type = "s"
test_path = Path("scratch/test_out.xlsx")
test_path.parent.mkdir(exist_ok=True)
wb.save(test_path)

wb2 = openpyxl.load_workbook(test_path, data_only=False)
ws2 = wb2.active
print("CELL A1:", ws2["A1"].data_type, ws2["A1"].value)
