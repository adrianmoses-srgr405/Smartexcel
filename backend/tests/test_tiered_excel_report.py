import pytest
from pathlib import Path
import pandas as pd
import openpyxl
from app.services.formula_planner import FormulaPlanner
from app.services.task_decomposer import TaskDecomposer
from app.services.report_service import ReportService
from app.services.hybrid_ai_service import HybridAIService

DATASET_PATH = Path("../storage/uploads/98a687aad2384547923c8ed474cd637f_Data_Dummy_Penjualan_Mobil_Pekanbaru_1000 (1).xlsx")

@pytest.fixture(scope="module")
def car_1000_df():
    assert DATASET_PATH.exists(), f"Dataset file not found at {DATASET_PATH.resolve()}"
    df = pd.read_excel(DATASET_PATH, sheet_name="Data_Penjualan")
    assert len(df) == 1000, f"Expected 1000 rows in Data_Penjualan, got {len(df)}"
    return df

def test_formula_planner_grouping_and_aggregations(car_1000_df):
    """
    Verifies that Formula Planner:
    1. Identifies Model as grouping dimension for 'setiap tipe mobil'
    2. Retrieves all 20 unique car models dynamically without hard-coding
    3. Detects relevant summable columns: Harga, Diskon, Harga_Netto, DP
    4. Excludes non-aggregation columns: ID, Tanggal, Tahun, CC, Tenor, Usia
    """
    query = "Rekap total penjualan setiap tipe mobil dari seluruh data."
    tasks = TaskDecomposer.decompose(query, sample_df=car_1000_df)
    assert len(tasks) >= 1
    assert tasks[0].get("group_by") == ["Model"]
    assert tasks[0].get("operation") == "SUM"

    plan = FormulaPlanner.plan_formula(query, tasks=tasks, sample_df=car_1000_df)
    assert plan["grouping"] == "Model"
    assert len(plan["groups"]) == 20
    # Verify dynamic car model names exist in plan
    expected_sample = ["Avanza", "Fortuner", "Innova", "Xpander", "Brio", "Rush"]
    for m in expected_sample:
        assert m in plan["groups"]

    # Verify aggregations
    agg_cols = [a["column"] for a in plan["aggregations"]]
    assert "Harga_Netto" in agg_cols
    assert "Harga" in agg_cols
    assert "Diskon" in agg_cols
    assert "DP" in agg_cols

    # Verify non-aggregations are NOT in aggregations
    excluded = ["ID_Transaksi", "Tanggal", "Tahun", "CC", "Tenor_Bulan", "Usia_Pembeli"]
    for ex in excluded:
        assert ex not in agg_cols

def test_query_variations_planning(car_1000_df):
    """
    Verifies that Formula Planner supports:
    1. 'rata-rata penjualan setiap tipe mobil' -> AVERAGE
    2. 'jumlah transaksi setiap tipe mobil' -> COUNT
    3. 'total harga dan jumlah unit setiap tipe mobil' -> SUM + COUNT
    """
    # 1. Rata-rata
    q1 = "rata-rata penjualan setiap tipe mobil"
    t1 = TaskDecomposer.decompose(q1, sample_df=car_1000_df)
    p1 = FormulaPlanner.plan_formula(q1, tasks=t1, sample_df=car_1000_df)
    assert p1["grouping"] == "Model"
    assert any(a["function"] == "AVERAGE" for a in p1["aggregations"])

    # 2. Jumlah transaksi
    q2 = "jumlah transaksi setiap tipe mobil"
    t2 = TaskDecomposer.decompose(q2, sample_df=car_1000_df)
    p2 = FormulaPlanner.plan_formula(q2, tasks=t2, sample_df=car_1000_df)
    assert p2["grouping"] == "Model"
    assert any(a["function"] == "COUNT" for a in p2["aggregations"])

    # 3. Total harga dan jumlah unit
    q3 = "total harga dan jumlah unit setiap tipe mobil"
    t3 = TaskDecomposer.decompose(q3, sample_df=car_1000_df)
    p3 = FormulaPlanner.plan_formula(q3, tasks=t3, sample_df=car_1000_df)
    assert p3["grouping"] == "Model"
    functions = [a["function"] for a in p3["aggregations"]]
    assert "SUM" in functions
    assert "COUNT" in functions

def test_tiered_excel_export_1000_rows(car_1000_df):
    """
    Verifies complete Excel generation from the 1000 raw rows:
    1. Output workbook contains 2 sheets ('Rekap Penjualan', 'Data Mentah')
    2. 'Data Mentah' contains all 1000 rows as source of truth
    3. 'Rekap Penjualan' contains all 20 unique car types with their details
    4. Each car type has a subtotal row with active =SUM(...) formula
    5. Final grand total row (TOTAL KESELURUHAN) exists with active =SUM(...) formula
    6. File can be reopened and validated cleanly with openpyxl
    """
    query = "Rekap total penjualan setiap tipe mobil dari seluruh data."
    tasks = TaskDecomposer.decompose(query, sample_df=car_1000_df)
    plan = FormulaPlanner.plan_formula(query, tasks=tasks, sample_df=car_1000_df)

    export_payload = {
        "user_query": query,
        "formula_id": "SUM",
        "formula_name": "SUM",
        "generated_excel_formula": "=SUMIF(Model, \"Avanza\", Harga_Netto)",
        "formula_explanation": "Rekap total penjualan bertingkat per tipe mobil",
        "tasks": tasks,
        "formula_plan": plan,
        "group_by": [plan["grouping"]],
        "raw_df": car_1000_df
    }

    export_path = ReportService.export_analysis_to_excel(export_payload, filename_prefix="Test_Laporan_Penjualan_Mobil")
    assert export_path.exists()

    # Reopen and inspect workbook
    wb = openpyxl.load_workbook(export_path, data_only=False)
    assert "Rekap Penjualan" in wb.sheetnames
    assert "Data Mentah" in wb.sheetnames

    # Verify Data Mentah
    ws_raw = wb["Data Mentah"]
    assert ws_raw.max_row == 1001 # 1 header + 1000 data rows

    # Verify Rekap Penjualan
    ws_rep = wb["Rekap Penjualan"]
    subtotals = []
    grand_total_found = False

    for r in range(1, ws_rep.max_row + 1):
        cell_val = str(ws_rep.cell(row=r, column=1).value or "")
        if cell_val.startswith("TOTAL ") and cell_val != "TOTAL KESELURUHAN":
            last_col_val = str(ws_rep.cell(row=r, column=ws_rep.max_column).value or "")
            subtotals.append((cell_val, last_col_val))
            # Verify formula syntax
            assert last_col_val.startswith("=SUM("), f"Expected =SUM formula in subtotal at row {r}, got {last_col_val}"
        elif cell_val == "TOTAL KESELURUHAN":
            grand_total_found = True
            gt_formula = str(ws_rep.cell(row=r, column=ws_rep.max_column).value or "")
            assert gt_formula.startswith("=SUM("), f"Expected =SUM formula in grand total, got {gt_formula}"

    assert len(subtotals) == 20, f"Expected 20 subtotals (1 per model), found {len(subtotals)}"
    assert grand_total_found, "TOTAL KESELURUHAN not found in report"

    # Specific car models check
    subtotal_labels = [s[0] for s in subtotals]
    assert "TOTAL AVANZA" in subtotal_labels
    assert "TOTAL FORTUNER" in subtotal_labels
    assert "TOTAL INNOVA" in subtotal_labels
    assert "TOTAL XPANDER" in subtotal_labels
