import pytest
import pandas as pd
from app.services.task_decomposer import TaskDecomposer
from app.services.hybrid_ai_service import HybridAIService
from app.services.report_service import ReportService
import openpyxl

@pytest.fixture
def car_columns_profile():
    return [
        {"original_name": "ID", "column_letter": "A", "inferred_type": "Text"},
        {"original_name": "Tanggal", "column_letter": "B", "inferred_type": "Date"},
        {"original_name": "Sales", "column_letter": "C", "inferred_type": "Text"},
        {"original_name": "Cabang", "column_letter": "D", "inferred_type": "Text"},
        {"original_name": "Kota", "column_letter": "E", "inferred_type": "Text"},
        {"original_name": "Merek", "column_letter": "F", "inferred_type": "Text", "sample_values": ["Toyota", "Honda"]},
        {"original_name": "Model", "column_letter": "G", "inferred_type": "Text", "sample_values": ["Avanza", "Innova", "Brio"]},
        {"original_name": "Tahun", "column_letter": "H", "inferred_type": "Numeric"},
        {"original_name": "Tipe", "column_letter": "I", "inferred_type": "Text"},
        {"original_name": "Segmen", "column_letter": "J", "inferred_type": "Text"},
        {"original_name": "CC", "column_letter": "K", "inferred_type": "Numeric"},
        {"original_name": "Transmisi", "column_letter": "L", "inferred_type": "Text"},
        {"original_name": "Bahan_Bakar", "column_letter": "M", "inferred_type": "Text"},
        {"original_name": "Warna", "column_letter": "N", "inferred_type": "Text"},
        {"original_name": "Harga", "column_letter": "O", "inferred_type": "Numeric"},
        {"original_name": "Diskon", "column_letter": "P", "inferred_type": "Numeric"},
        {"original_name": "Harga_Netto", "column_letter": "Q", "inferred_type": "Numeric"},
        {"original_name": "Metode_Bayar", "column_letter": "R", "inferred_type": "Text"},
        {"original_name": "Tenor", "column_letter": "S", "inferred_type": "Numeric"},
        {"original_name": "DP", "column_letter": "T", "inferred_type": "Numeric"},
    ]

@pytest.fixture
def car_df():
    return pd.DataFrame({
        "Merek": ["Toyota", "Toyota", "Honda", "Toyota"],
        "Model": ["Avanza", "Avanza", "Brio", "Innova"],
        "Tahun": [2024, 2024, 2023, 2024],
        "Harga": [250000000, 260000000, 180000000, 380000000],
        "Harga_Netto": [240000000, 250000000, 175000000, 370000000],
        "DP": [50000000, 52000000, 36000000, 76000000],
    })

def test_1_total_harga_avanza(car_columns_profile):
    """Test 1: 'Total harga Avanza' -> SUMIF(G:G, 'Avanza', O:O)"""
    tasks = TaskDecomposer.decompose("Total harga Avanza", columns_profile=car_columns_profile)
    assert len(tasks) == 1
    t = tasks[0]
    res = HybridAIService._process_single_task(
        task_context=t,
        raw_query="Total harga Avanza",
        columns_profile=car_columns_profile,
        schema_col_names=[c["original_name"] for c in car_columns_profile]
    )
    assert res["formula_name"] == "SUMIF"
    assert res["generated_formula"] == '=SUMIF(G:G, "Avanza", O:O)'

def test_2_total_harga_netto_avanza(car_columns_profile):
    """Test 2: 'Total harga netto Avanza' -> SUMIF(G:G, 'Avanza', Q:Q)"""
    tasks = TaskDecomposer.decompose("Total harga netto Avanza", columns_profile=car_columns_profile)
    assert len(tasks) == 1
    t = tasks[0]
    res = HybridAIService._process_single_task(
        task_context=t,
        raw_query="Total harga netto Avanza",
        columns_profile=car_columns_profile,
        schema_col_names=[c["original_name"] for c in car_columns_profile]
    )
    assert res["formula_name"] == "SUMIF"
    assert res["generated_formula"] == '=SUMIF(G:G, "Avanza", Q:Q)'

def test_3_total_dp_avanza(car_columns_profile):
    """Test 3: 'Total DP Avanza' -> SUMIF(G:G, 'Avanza', T:T)"""
    tasks = TaskDecomposer.decompose("Total DP Avanza", columns_profile=car_columns_profile)
    assert len(tasks) == 1
    t = tasks[0]
    res = HybridAIService._process_single_task(
        task_context=t,
        raw_query="Total DP Avanza",
        columns_profile=car_columns_profile,
        schema_col_names=[c["original_name"] for c in car_columns_profile]
    )
    assert res["formula_name"] == "SUMIF"
    assert res["generated_formula"] == '=SUMIF(G:G, "Avanza", T:T)'

def test_4_total_harga_harga_netto_dan_dp_avanza(car_columns_profile):
    """Test 4: 'Total harga, harga netto dan DP Avanza' -> 3 calculation tasks"""
    tasks = TaskDecomposer.decompose("Total harga, harga netto dan DP Avanza", columns_profile=car_columns_profile)
    assert len(tasks) == 3
    targets = [t["target_term"] for t in tasks]
    assert "Harga" in targets
    assert "Harga_Netto" in targets
    assert "DP" in targets
    for t in tasks:
        assert t["operation"] == "SUM"
        assert t["filters"][0]["field"] == "Model"
        assert t["filters"][0]["value"] == "Avanza"

def test_5_rangkap_data_dan_total_harga_netto_dp(car_columns_profile):
    """Test 5: 'Rangkap data mobil Avanza dan total harga, harga netto dan DP' -> 1 retrieval + 3 calculation tasks"""
    query = "Rangkap data mobil Avanza dan total harga, harga netto dan DP"
    tasks = TaskDecomposer.decompose(query, columns_profile=car_columns_profile)
    assert len(tasks) == 4
    
    assert tasks[0]["task_type"] == "retrieval"
    assert tasks[0]["operation"] == "FILTER"
    assert tasks[0]["filters"][0]["field"] == "Model"
    assert tasks[0]["filters"][0]["value"] == "Avanza"

    calc_targets = [t["target_term"] for t in tasks[1:]]
    assert calc_targets == ["Harga", "Harga_Netto", "DP"]
    for t in tasks[1:]:
        assert t["task_type"] == "calculation"
        assert t["operation"] == "SUM"
        assert t["filters"][0]["field"] == "Model"
        assert t["filters"][0]["value"] == "Avanza"

    # Process all 4 tasks
    processed = []
    for t in tasks:
        r = HybridAIService._process_single_task(
            task_context=t,
            raw_query=query,
            columns_profile=car_columns_profile,
            schema_col_names=[c["original_name"] for c in car_columns_profile]
        )
        processed.append(r)

    assert processed[0]["formula"] == '=FILTER(A:Z, G:G="Avanza")'
    assert processed[1]["formula"] == '=SUMIF(G:G, "Avanza", O:O)'
    assert processed[2]["formula"] == '=SUMIF(G:G, "Avanza", Q:Q)'
    assert processed[3]["formula"] == '=SUMIF(G:G, "Avanza", T:T)'

def test_6_total_harga_toyota_avanza_tahun_2024(car_columns_profile):
    """Test 6: 'Total harga Toyota Avanza tahun 2024' -> SUMIFS with Merek=Toyota, Model=Avanza, Tahun=2024"""
    query = "Total harga Toyota Avanza tahun 2024"
    tasks = TaskDecomposer.decompose(query, columns_profile=car_columns_profile)
    assert len(tasks) == 1
    t = tasks[0]
    res = HybridAIService._process_single_task(
        task_context=t,
        raw_query=query,
        columns_profile=car_columns_profile,
        schema_col_names=[c["original_name"] for c in car_columns_profile]
    )
    assert res["formula_name"] == "SUMIFS"
    assert res["generated_formula"] == '=SUMIFS(O:O, F:F, "Toyota", G:G, "Avanza", H:H, "2024")'

def test_7_real_data_execution_and_excel_output(car_df):
    """Test real query execution on car dataframe and verify Excel output file generation"""
    query = "Rangkap data mobil Avanza dan total kan harga nya, harga netto dan dp"
    res = HybridAIService.process_query(query, sample_df=car_df)
    
    tasks = res["tasks"]
    assert len(tasks) == 4
    
    # Task 1: Retrieval
    assert tasks[0]["task_type"] == "retrieval"
    assert tasks[0]["operation"] == "FILTER"
    assert tasks[0]["result"] == 2  # 2 matching Avanza rows in car_df

    # Task 2: SUM Harga
    assert tasks[1]["target"] == "Harga"
    assert tasks[1]["result"] == 510000000

    # Task 3: SUM Harga_Netto
    assert tasks[2]["target"] == "Harga_Netto"
    assert tasks[2]["result"] == 490000000

    # Task 4: SUM DP
    assert tasks[3]["target"] == "DP"
    assert tasks[3]["result"] == 102000000

    # Test Excel Export
    export_payload = {
        "user_query": query,
        "formula_id": res["formula_name"],
        "formula_name": res["formula_name"],
        "generated_excel_formula": res["generated_formula"],
        "formula_explanation": res["explanation"]["formula_selection"],
        "tasks": tasks,
        "table_headers": list(car_df.columns),
        "table_rows": car_df.to_dict(orient="records"),
        "calculation_summary": {"target_field": "Harga", "grand_total": 510000000}
    }
    export_path = ReportService.export_analysis_to_excel(export_payload, filename_prefix="Test_Car_Export")
    assert export_path.exists()

    # Reopen and inspect cells
    wb = openpyxl.load_workbook(export_path)
    ws = wb.active
    assert ws.title == "Ringkasan Laporan"
    assert ws["A1"].value == "SISTEM PEMODELAN LAPORAN & FORMULA EXCEL OTOMATIS - PTPN"
    # Multi-task headers at row 10
    assert ws.cell(10, 2).value == "Task ID"
    assert ws.cell(11, 2).value == "task_1"
    assert ws.cell(12, 2).value == "task_2"
    assert ws.cell(13, 2).value == "task_3"
    assert ws.cell(14, 2).value == "task_4"
