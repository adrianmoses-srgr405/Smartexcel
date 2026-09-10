import pytest
import pandas as pd
import openpyxl
from pathlib import Path

from app.services.grouping_subtotal_service import GroupingSubtotalService
from app.services.execution_engine import ExecutionEngine
from app.services.report_service import ReportService
from app.schemas.intent import StructuredAnalysisIntent, FormulaDecisionResult

@pytest.fixture
def sample_car_df():
    # Miniature dataset mimicking Data_Dummy_Penjualan_Mobil_Pekanbaru_1000
    data = [
        {"ID_Transaksi": "TRX001", "Sales": "Andi", "Cabang": "Pekanbaru Panam", "Merek": "Wuling", "Model": "Almaz", "Tahun": 2024, "Tipe": "V", "Harga": 380000000, "Diskon": 10000000, "Harga_Netto": 370000000, "DP": 70000000},
        {"ID_Transaksi": "TRX002", "Sales": "Bagus", "Cabang": "Pekanbaru Sudirman", "Merek": "Wuling", "Model": "Almaz", "Tahun": 2024, "Tipe": "G", "Harga": 350000000, "Diskon": 5000000, "Harga_Netto": 345000000, "DP": 60000000},
        {"ID_Transaksi": "TRX003", "Sales": "Sari", "Cabang": "Pekanbaru Panam", "Merek": "Wuling", "Model": "Alvez", "Tahun": 2024, "Tipe": "EX", "Harga": 220000000, "Diskon": 5000000, "Harga_Netto": 215000000, "DP": 40000000},
        {"ID_Transaksi": "TRX004", "Sales": "Yoga", "Cabang": "Pekanbaru Harapan Raya", "Merek": "Wuling", "Model": "Alvez", "Tahun": 2024, "Tipe": "CE", "Harga": 200000000, "Diskon": 4000000, "Harga_Netto": 196000000, "DP": 35000000},
    ]
    return pd.DataFrame(data)

def test_detect_grouping_column_explicit(sample_car_df):
    # Query mentions 'per cabang' -> Cabang
    grp_col = GroupingSubtotalService.detect_grouping_column(sample_car_df, query="total penjualan per cabang")
    assert grp_col == "Cabang"

    # Query mentions 'per sales' -> Sales
    grp_col = GroupingSubtotalService.detect_grouping_column(sample_car_df, query="rekap data per sales")
    assert grp_col == "Sales"

    # Query mentions 'tipe mobil' / 'model' -> Model
    grp_col = GroupingSubtotalService.detect_grouping_column(sample_car_df, query="rekap total penjualan setiap tipe mobil dari seluruh data")
    assert grp_col in ["Model", "Tipe"]

def test_detect_grouping_column_structural(sample_car_df):
    # General query 'total' without explicit dimension -> chooses best hierarchy (Model or Cabang)
    grp_col = GroupingSubtotalService.detect_grouping_column(sample_car_df, query="hitung total")
    assert grp_col in ["Model", "Cabang", "Sales", "Merek"]

def test_detect_measure_columns(sample_car_df):
    # Target 'total harga' -> ['Harga'] (strictly excludes 'Harga_Netto' and 'Tahun')
    measures = GroupingSubtotalService.detect_measure_columns(sample_car_df, query="total harga")
    assert measures == ["Harga"]
    assert "Tahun" not in measures

    # Target 'total harga dan dp' -> ['Harga', 'DP']
    measures = GroupingSubtotalService.detect_measure_columns(sample_car_df, query="total harga dan dp")
    assert "Harga" in measures
    assert "DP" in measures
    assert "Tahun" not in measures

    # General 'total' -> all numeric measures (Harga, Diskon, Harga_Netto, DP), never Tahun
    measures = GroupingSubtotalService.detect_measure_columns(sample_car_df, query="total")
    assert "Harga" in measures
    assert "DP" in measures
    assert "Diskon" in measures
    assert "Harga_Netto" in measures
    assert "Tahun" not in measures
    assert "ID_Transaksi" not in measures

def test_build_inline_subtotals(sample_car_df):
    rows, headers, meta = GroupingSubtotalService.build_inline_subtotals(
        sample_car_df,
        grouping_col="Model",
        measure_cols=["Harga", "DP"],
        query="total"
    )

    # 4 data rows + 2 subtotal rows (Almaz, Alvez) + 1 grand total row = 7 rows total
    assert len(rows) == 7
    assert meta["has_subtotals"] is True
    assert meta["total_groups"] == 2
    assert "Almaz" in meta["groups"]
    assert "Alvez" in meta["groups"]

    # Check that Almaz subtotal appears after Almaz detail rows
    almaz_rows = [r for r in rows if r.get("_group_val") == "Almaz"]
    assert len(almaz_rows) == 3 # 2 data + 1 subtotal
    assert almaz_rows[-1]["_is_subtotal"] is True
    assert almaz_rows[-1]["Model"] == "TOTAL"
    assert almaz_rows[-1]["Harga"] == 380000000 + 350000000 # 730,000,000

    # Check Alvez subtotal
    alvez_rows = [r for r in rows if r.get("_group_val") == "Alvez"]
    assert len(alvez_rows) == 3 # 2 data + 1 subtotal
    assert alvez_rows[-1]["_is_subtotal"] is True
    assert alvez_rows[-1]["Model"] == "TOTAL"
    assert alvez_rows[-1]["Harga"] == 220000000 + 200000000 # 420,000,000

    # Check Grand Total row
    grand_total_row = rows[-1]
    assert grand_total_row["_is_grand_total"] is True
    assert grand_total_row["Model"] == "TOTAL KESELURUHAN"
    assert grand_total_row["Harga"] == 730000000 + 420000000 # 1,150,000,000

def test_execution_engine_with_dynamic_subtotals(sample_car_df):
    intent = StructuredAnalysisIntent(
        intent="aggregation",
        operation="SUM",
        target_field="Harga",
        confidence_score=0.95,
        user_explanation="Rekapitulasi penjualan mobil"
    )
    decision = FormulaDecisionResult(
        formula_id="SUM",
        formula_name="SUM",
        category="math",
        reason="Aggregation request",
        syntax_pattern="=SUM(range)",
        generated_excel_formula="=SUM(H2:H5)",
        is_confident=True
    )

    summary, headers, rows, chart_data = ExecutionEngine.execute(
        df=sample_car_df,
        intent=intent,
        decision=decision,
        user_query="rekap total penjualan setiap tipe mobil dari seluruh data"
    )

    assert summary.get("has_subtotals") is True
    assert summary.get("grouping_column") in ["Model", "Tipe"]
    assert len(rows) > len(sample_car_df) # contains inline subtotal rows
    assert any(r.get("_is_subtotal") for r in rows)
    assert len(chart_data) >= 2 # Chart data per model

def test_excel_export_formulas_for_subtotals(sample_car_df, tmp_path):
    rows, headers, meta = GroupingSubtotalService.build_inline_subtotals(
        sample_car_df,
        grouping_col="Model",
        measure_cols=["Harga", "DP"],
        query="total"
    )

    analysis_data = {
        "user_query": "rekap total penjualan setiap tipe mobil dari seluruh data",
        "formula_id": "SUM",
        "formula_name": "SUM",
        "tasks": []
    }

    export_path = ReportService._export_displayed_table(
        analysis_data=analysis_data,
        filename_prefix="test_subtotal_export",
        table_headers=headers,
        table_rows=rows,
        raw_df=sample_car_df
    )

    assert Path(export_path).exists()

    # Load workbook with openpyxl and inspect formulas
    wb = openpyxl.load_workbook(export_path, data_only=False)
    ws = wb["Rekap Penjualan"]

    # Find Harga column index
    harga_col_letter = None
    tahun_col_letter = None
    for col_idx in range(1, ws.max_column + 1):
        header_val = str(ws.cell(row=4, column=col_idx).value or "").upper()
        if header_val == "HARGA":
            harga_col_letter = openpyxl.utils.get_column_letter(col_idx)
        elif header_val == "TAHUN":
            tahun_col_letter = openpyxl.utils.get_column_letter(col_idx)

    assert harga_col_letter is not None
    assert tahun_col_letter is not None

    # Inspect all rows for formulas
    formula_cells = []
    for r in range(5, ws.max_row + 1):
        c_val = str(ws[f"{harga_col_letter}{r}"].value or "")
        t_val = str(ws[f"{tahun_col_letter}{r}"].value or "")
        # Tahun must NEVER have a formula
        assert not t_val.startswith("="), f"Tahun column at row {r} should not contain formula: {t_val}"

        if c_val.startswith("="):
            formula_cells.append((r, c_val))

    # We should have subtotal formulas for Almaz, Alvez, and Grand Total
    assert len(formula_cells) >= 3
    # Subtotal 1 (Almaz): =SUM(Harga5:Harga6)
    assert "=SUM(" in formula_cells[0][1]
    # Subtotal 2 (Alvez): =SUM(Harga8:Harga9)
    assert "=SUM(" in formula_cells[1][1]
    # Grand Total: =SUM(Harga7,Harga10) or =SUM(range)
    assert "=SUM(" in formula_cells[2][1]
